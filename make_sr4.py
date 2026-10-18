from docx import Document
from docx.shared import Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
src='D:/xuan/英语/ESS_101_SR3_Group19_Shirley_25722118.docx'
d=Document(src)
d.paragraphs[1].text='Article Number: 4'
d.paragraphs[3].text='Group:19\tWeek 4\t\tESS101'
info=d.tables[0]
def put(c,txt):
 c.text=txt
 for p in c.paragraphs:
  p.paragraph_format.space_after=Pt(5)
  p.paragraph_format.line_spacing=1.08
  for r in p.runs:r.font.name='Calibri';r.font.size=Pt(11);r.font.color.rgb=RGBColor(0,0,0)
vals={0:'Kannan, R.; Firke, R.; Bengle, S.; Deshmukh, S.',1:'GreenBench: Benchmarking Energy Efficiency and Carbon Footprint of Open-Source LLM Inference on Apple Silicon',2:'Not a journal article. arXiv preprint; the authors report acceptance at IEEE ICCUBEA 2026.',3:'2026',4:'Not applicable',5:'https://arxiv.org/abs/2608.28667v1',6:'Kannan, R., Firke, R., Bengle, S., & Deshmukh, S. (2026). GreenBench: Benchmarking energy efficiency and carbon footprint of open-source LLM inference on Apple Silicon. arXiv. https://doi.org/10.48550/arXiv.2608.28667 (Version 1, 24 August 2026; accessed 15 October 2026).'}
for i,v in vals.items():put(info.rows[i].cells[1],v)
put(info.rows[3].cells[3],'August (arXiv submission)');put(info.rows[4].cells[3],'Not applicable')
answers=[
'''1. This article connects computer science with my essay on sustainability. It moves the discussion from the general idea of Green AI to choosing models for everyday use.
2. It helps me distinguish training from inference. A model continues to consume energy each time someone uses it, so sustainable deployment matters as well as sustainable development.
3. It also encourages me to ask what an efficiency figure actually measures. A low number is not convincing unless the measurement boundary and assumptions are clear.''',
'''1. A strength is the comparison of several models under the same hardware and quantization settings. The tables make the relationship between answer quality, speed and energy easier to evaluate.
2. However, system power is estimated rather than measured at the wall socket. A proper follow-up should use a calibrated power meter and repeat tests on several machines.
3. The carbon units need checking. Equation 3, 0.09 J/token and 708 g CO2/kWh imply about 17.7 mg CO2 per 1,000 tokens; Table IV gives 0.18 mg. This appears to be a roughly 100-fold discrepancy. I would not quote that table's carbon figures without clarification.
4. Code output is not tested for correctness. An efficient wrong answer may require another attempt, so energy per successful task would be more informative than energy per token alone.''',
'''1. I would use the comparison between model size, energy and answer quality in the Solution section of my SPSE essay. It illustrates how selecting a suitable model could avoid unnecessary computation.
2. For Evaluation, I would discuss the restricted hardware setting and estimated whole-system power. The study cannot establish that laptops are always greener than cloud services.
3. I would link this reading to Green AI: efficiency should accompany accuracy. I would also retain the warning from Sustainable Computing - Without the Hot Air that savings per task do not guarantee lower total demand.''',
'''1. The version examined is arXiv v1, submitted on 24 August 2026, so it is recent relative to this SR's date, 15 October 2026.
2. The reference should identify the version actually read. An arXiv submission date is not necessarily the date of a final conference publication.
3. Its results concern the tested models and software. New models, inference engines or hardware may change the comparison; publication recency alone does not make a result generally applicable.''',
'''1. The paper provides named authors, affiliations, methods and references. Kannan is affiliated with Pimpri Chinchwad College of Engineering, Firke with Red Hat, and Bengle and Deshmukh are listed as independent researchers.
2. The arXiv record says the paper was accepted at IEEE ICCUBEA 2026. I am citing the accessible arXiv version; I have not independently verified the final proceedings record.
3. Repository availability and conference acceptance are useful context, but they do not replace checking measurements and calculations. The apparent carbon-unit inconsistency lowers my confidence in those figures.
4. I would use this source as a focused case study and compare its claims with other research, rather than treating it as definitive evidence about all AI systems.''',
'''Article summary
Purpose: To compare the energy and performance of local language-model inference on consumer hardware.
Method: Five models, three task types and over 1,500 inference records on one M4 Pro laptop with 48 GB memory. Package power readings are combined with timing logs; whole-system power is estimated.
Reported findings:
1. Smaller models use up to 62% less energy per token in the tested setup.
2. Qwen 2.5 scores 57% on the selected MMLU questions, compared with Gemma 2's 61%, while offering higher throughput.
3. The authors recommend reporting energy per token alongside performance measures.
My evaluation: These findings support testing model choices, but they do not prove universal environmental savings. Hardware coverage, answer quality, electricity assumptions and the apparent carbon-unit error limit the conclusions.
Essay connection: Green AI is a possible solution; reliable measurement and task-specific evaluation determine how convincing that solution is.'''
]
t=d.tables[1]
for i,a in enumerate(answers):put(t.rows[i].cells[1 if i<5 else 0],a)
vocab=[
('inference','n.','/ˈɪnfərəns/','推理','using a trained model to produce an output','infer (v.); inferential (adj.)','run inference; inference cost','prediction (related term); training (contrast)','I Introduction','Inference on a phone may continue without an internet connection.','We should compare the energy needed for one useful inference.'),
('throughput','n.','/ˈθruːpʊt/','吞吐量','the amount of work completed in a unit of time','No common derived form','high throughput; inference throughput','processing rate; low throughput (contrast)','IV-A Inference Throughput Analysis','Higher throughput means more tokens can be produced each second.','Batching may improve throughput while increasing waiting time.'),
('latency','n.','/ˈleɪtənsi/','延迟','the time spent waiting for a response or operation','latent (adj., broader meaning)','low latency; reduce latency','delay; immediacy (contrast)','III-D2 Timing Instrumentation','A voice assistant needs low latency to feel responsive.','Low energy use does not automatically guarantee low latency.'),
('quantization','n.','/ˌkwɒntaɪˈzeɪʃən/','量化','representing numerical values using a limited set of levels','quantize (v.); quantized (adj.)','four-bit quantization; quantization level','reduced precision (related term); full precision (contrast)','II-C Quantization and Model Compression','Quantization can help a model fit into limited memory.','We should check whether quantization changes answer quality.'),
('granularity','n.','/ˌɡrænjʊˈlærəti/','精细程度','the level of detail at which information is measured','granular (adj.)','measurement granularity; fine granularity','level of detail; coarseness (contrast)','V Limitations and Future Work','Finer measurement granularity can reveal short power spikes.','The report explains the granularity of its electricity data.'),
('reproducibility','n.','/ˌriːprəˌdjuːsəˈbɪləti/','可复现性','the extent to which results can be obtained again using a stated method','reproduce (v.); reproducible (adj.)','ensure reproducibility; reproducibility of results','repeatability (related term); irreproducibility (opposite)','III-E Experimental Protocol','Publishing experimental settings improves reproducibility.','A separate laboratory should check the reproducibility of the findings.'),
('trade-off','n.','/ˈtreɪd ɒf/','权衡','a situation where gaining one benefit involves giving up another','trade off (phrasal v.)','accuracy-energy trade-off; balance a trade-off','compromise; no exact opposite','IV-C Accuracy-Efficiency Pareto Analysis','Choosing a smaller model may involve a trade-off between speed and quality.','The best trade-off depends on the task.'),
('deployment','n.','/dɪˈplɔɪmənt/','部署','putting a system into practical use','deploy (v.); deployed (adj.)','local deployment; model deployment','implementation (related term); withdrawal (contrast)','IV-E Implications for Apple Silicon Deployment','Local deployment keeps model execution on the user\'s device.','Deployment decisions should consider security as well as energy.'),
('contention','n.','/kənˈtenʃən/','资源争用','competition between processes for a shared computing resource','contend (v.); contending (adj.)','memory contention; resource contention','competition; uncontested access (contrast)','III-E Experimental Protocol','Running several models together may cause memory contention.','The experiment isolates models to reduce resource contention.'),
('generalization','n.','/ˌdʒenrəlaɪˈzeɪʃən/','推广适用性','extending a finding beyond the particular cases examined','generalize (v.); general (adj.)','generalization across devices; limit generalization','extension; case-specific conclusion (contrast)','V Limitations and Future Work','Results from one laptop allow only limited generalization.','Testing several devices would support broader generalization.')]
entries=['Vocabulary development\nArticle-context locations below identify where the term is used; they are not invented quotations. Context II sentences are original practice examples.']
for n,(word,pos,ipa,cn,definition,family,coll,syn,loc,a,b) in enumerate(vocab,1):
 entries.append(f'{n}. Word: {word} ({pos}) {ipa} {cn}\nDefinition: {definition}.\nWord family: {family}.\nCollocations: {coll}.\nSynonyms / Antonyms: {syn}.\nWord in Context I: See {loc} in Kannan et al. (2026), where this term is used in discussing the study.\nWord in Context II: (a) {a} (b) {b}')
put(t.rows[6].cells[0],'\n\n'.join(entries))
# Preserve the supplied form but remove fixed row heights and inherited keep rules.
for table in d.tables:
 for row in table.rows:
  for x in list(row._tr.trPr or []):
   if x.tag==qn('w:trHeight'):row._tr.trPr.remove(x)
  for cell in row.cells:
   for p in cell.paragraphs:p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=False
for sec in d.sections:
 sec.header.paragraphs[0].text=''
out=Path('output/docx');out.mkdir(parents=True,exist_ok=True)
f=out/'ESS_101_SR4_Group19_Shirley_25722118.docx';d.save(f);print(f.resolve())
